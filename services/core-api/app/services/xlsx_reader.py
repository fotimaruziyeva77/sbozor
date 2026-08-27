"""Xavfsiz `.xlsx` o'qish — IKKITA MUSTAQIL DoS YUZASI (T-02-87, T-02-88).

=============================================================================
TARTIB MAJBURIY VA U SHU YERDA LITERAL YOZILGAN.

`.xlsx` — bu ZIP, ZIP ichida esa XML. Ikkala qatlam ham mustaqil hujum
yuzasi va ULAR UCHUN HIMOYA HAM MUSTAQIL:

  1. XML  — billion laughs / quadratic blowup. Himoya: `defusedxml`.
            openpyxl O'Z HUJJATIDA buni talab qiladi: standart holatda u
            bu hujumlardan HIMOYALANMAGAN.
  2. ZIP  — zip-bomba (1 MB fayl 10 GB ga ochiladi). Himoya: OCHILGAN
            hajm chegarasi.

⚠ `defusedxml` ZIP-BOMBAGA HECH QANDAY YORDAM BERMAYDI. U faqat XML
parseriga tegadi, zip-bomba esa parserga YETIB BORMASDAN, `zipfile`
ochayotgan paytda xotirani yeydi. Shuning uchun ZIP darvozasi
`load_workbook()` dan OLDIN turadi va `test_zip_bomb_is_rejected_before_
parsing` aynan shu TARTIBNI (chaqirilgan/chaqirilmagan faktini) o'lchaydi
— natijani emas.
=============================================================================

DEFUSEDXML QANDAY ULANADI — O'LCHANGAN, TAXMIN QILINMAGAN:

openpyxl `defusedxml` ni O'RNATILGANLIGI bo'yicha aniqlaydi (o'z
`__init__.py` sida `import defusedxml` ni sinaydi va `openpyxl.DEFUSEDXML`
bayrog'ini qo'yadi). Konteynerda o'lchandi (openpyxl 3.1.5 + defusedxml
0.7.1 + Python 3.13):

    openpyxl.DEFUSEDXML                          -> True
    openpyxl.xml.functions.iterparse.__module__  -> defusedxml.common
    openpyxl.xml.functions.fromstring.__module__ -> defusedxml.common

Ya'ni openpyxl ning O'Z yo'li (`iterparse` HAM, `fromstring` HAM) to'liq
qamralgan va u bu modulning import TARTIBIGA bog'liq EMAS. Shuning uchun
haqiqiy darvoza — paketning o'rnatilganligi, va uni
`test_openpyxl_detected_defusedxml` AYNAN shu bayroq bo'yicha qulflaydi
(A9 ning smoke-testi). Bayroq `False` bo'lib qolsa (paket yo'qolsa yoki
openpyxl aniqlash mexanizmini o'zgartirsa) himoya JIMGINA yo'qolardi —
fayllar baribir o'qilaverardi.

`defuse_stdlib()` esa QO'SHIMCHA qatlam: u standart kutubxonaning XML
modullarini (`xml.etree.ElementTree.XMLParser`, `fromstring`, ...)
almashtiradi, ya'ni jarayonda openpyxl'dan TASHQARI XML parse qiladigan
har qanday yo'lni ham yopadi. openpyxl uchun u ZARUR EMAS (yuqoridagi
o'lchov), lekin zarari ham yo'q va u bu servisdagi yagona XML kirish
nuqtasi bo'lmay qolgan kunni oldindan qoplaydi.

-----------------------------------------------------------------------------
QATOR RAQAMLARI EXCEL BILAN AYNAN BIR XIL (D-14).

Foydalanuvchi xatoni Excelda tuzatadi, ya'ni "88-qator" xabari uning
ekranidagi 88-qatorga TO'G'RI KELISHI shart. Shuning uchun sarlavha
qatori TASHLANADI, lekin uning raqami (1) hisobdan CHIQMAYDI: birinchi
ma'lumot qatori — 2, uchinchisi — 4. Nol yoki birdan qayta sanash
xabarni bitta-ikkita qatorga surib, admin butunlay boshqa qatorni
tuzatishga urinardi.
-----------------------------------------------------------------------------

SARLAVHA MATNI UMUMAN O'QILMAYDI (O-05). Ustunlar FAQAT POZITSIYA
bo'yicha olinadi. Shablon foydalanuvchi tilida hosil bo'ladi
(`xlsx_template.build_template`), ya'ni ruscha shablonni yuklab olgan
admin uzbekcha interfeysda import qilsa ham fayl ishlaydi. Sarlavhaga
qarab ustun topish bu oqimni JIMGINA buzardi.
"""

from __future__ import annotations

import io
import warnings
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from typing import TYPE_CHECKING, Any, Final

import defusedxml
import openpyxl
import structlog
from defusedxml.ElementTree import ParseError

if TYPE_CHECKING:
    from collections.abc import Iterator

    from app.settings import Settings

__all__ = [
    "HEADER_ROW",
    "MAX_COLS",
    "MAX_ROWS",
    "MAX_SHEETS",
    "MAX_UNCOMPRESSED_BYTES",
    "MAX_UPLOAD_BYTES",
    "MAX_ZIP_ENTRIES",
    "ImportRejected",
    "ReadLimits",
    "SheetRow",
    "read_rows",
]

log = structlog.get_logger(__name__)

# Standart kutubxona XML modullarini almashtiradi — sabab modul
# docstringida ("QO'SHIMCHA qatlam" bandi). openpyxl uchun ZARUR EMAS:
# u `defusedxml` ni O'ZI aniqlaydi va bu chaqiruvga bog'liq emas.
#
# ⚠ OGOHLANTIRISH SO'NDIRILISHI — BIZNING SHOVQINIMIZ EMAS, UPSTREAM XATOSI.
# `defuse_stdlib()` ichida `with warnings.catch_warnings(): from . import
# cElementTree` yozilgan, lekin `simplefilter("ignore")` TUSHIB QOLGAN
# (defusedxml 0.7.1, manba o'qildi) — natijada har import'da
# `DeprecationWarning: defusedxml.cElementTree is deprecated` chiqadi.
# Uni so'ndirmaslik jurnalga har ishga tushishda YOLG'ON signal qo'shardi
# va haqiqiy ogohlantirishlar orasida ko'rinmay ketardi. So'ndirish
# AYNAN shu chaqiruv atrofida va AYNAN shu toifada — global filtr
# QO'YILMAYDI. (Bu A9 ning "paket 2021-yilgi" xavfining aniq ko'rinishi.)
with warnings.catch_warnings():
    warnings.simplefilter("ignore", DeprecationWarning)
    defusedxml.defuse_stdlib()

MAX_UPLOAD_BYTES: Final = 5 * 1024 * 1024
"""Tarmoqdan kelgan XOM `.xlsx` (ZIP) hajmi — 5 MB."""

MAX_UNCOMPRESSED_BYTES: Final = 50 * 1024 * 1024
"""ZIP OCHILGANDAGI umumiy hajm — 50 MB. `MAX_UPLOAD_BYTES` dan MUSTAQIL."""

MAX_ROWS: Final = 5_000
MAX_COLS: Final = 32
MAX_SHEETS: Final = 8
MAX_ZIP_ENTRIES: Final = 200
"""ZIP a'zolari soni — millionlab mayya faylli arxiv metama'lumot bilan
xotirani yeydi, garchi ochilgan hajm chegaradan past bo'lsa ham."""

HEADER_ROW: Final = 1
"""Excel qator raqami bo'yicha sarlavha qatori. Ma'lumot 2-qatordan boshlanadi."""


class ImportRejected(Exception):
    """Fayl PARSE QILINISHIDAN oldin (yoki paytida) rad etildi.

    `code` — `MARKET_ERROR_CODES` dagi qiymat (`file_too_large` /
    `file_too_complex` / `unsupported_file_type`). Router uni HTTP
    kodiga aylantiradi; xom istisno matni javobga HECH QACHON tushmaydi
    (T-02-58 bilan bir xil qoida — u kutubxona ichki tafsilotini
    oshkor qilardi).
    """

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True, slots=True)
class SheetRow:
    """Bitta ma'lumot qatori.

    `row` — EXCEL qator raqami (sarlavha = 1), ro'yxatdagi indeks EMAS.
    """

    row: int
    values: tuple[str | None, ...]


@dataclass(frozen=True, slots=True)
class ReadLimits:
    """O'qish chegaralari (A7 — SOZLANADIGAN).

    Standart qiymatlar modul konstantalaridan keladi;
    `ReadLimits.from_settings()` esa ularni `app/settings.py` orqali
    (ya'ni muhit o'zgaruvchisi bilan) qayta yozadi. A7 chegaralarni
    ochiq taxmin deb belgilaydi: 1000 rastadan kattaroq bozor kelganda
    ularga urilinadi.

    Testlar chegaralarni TORAYTIRIB beradi (masalan `rows=3`) — ya'ni
    darvoza mexanikasi 5000 qatorli fayl qurmasdan o'lchanadi. Standart
    QIYMATLARNING o'zi esa alohida test bilan qulflangan
    (`test_default_limits_match_the_documented_values`), aks holda
    kimdir `MAX_ROWS` ni 5 000 000 qilib qo'yganda toraytirilgan test
    baribir yashil qolardi.
    """

    upload_bytes: int = MAX_UPLOAD_BYTES
    uncompressed_bytes: int = MAX_UNCOMPRESSED_BYTES
    rows: int = MAX_ROWS
    cols: int = MAX_COLS
    sheets: int = MAX_SHEETS
    zip_entries: int = MAX_ZIP_ENTRIES

    @classmethod
    def from_settings(cls, settings: Settings) -> ReadLimits:
        """Ish vaqti sozlamalaridan quradi (12-faktor)."""
        return cls(
            upload_bytes=settings.import_max_upload_bytes,
            uncompressed_bytes=settings.import_max_uncompressed_bytes,
            rows=settings.import_max_rows,
            cols=settings.import_max_cols,
            sheets=settings.import_max_sheets,
            zip_entries=settings.import_max_zip_entries,
        )


_FILE_TOO_LARGE = "file_too_large"
_FILE_TOO_COMPLEX = "file_too_complex"
_UNSUPPORTED = "unsupported_file_type"

_PARSE_FAILURES: Final = (ValueError, ParseError, KeyError, TypeError, zipfile.BadZipFile)
"""Buzuq yoki hujumkor faylni o'qishda kutiladigan istisnolar.

`ValueError` ATAYIN keng: `defusedxml.common.DefusedXmlException` (ya'ni
`EntitiesForbidden`/`DTDForbidden`/`ExternalReferenceForbidden` uchalasi
ham) undan MEROS OLADI, openpyxl esa ularni O'Z `ValueError` iga o'raydi.
Konteynerda o'lchandi: billion-laughs `sharedStrings.xml`, `sheet1.xml`
va `workbook.xml` ning uchalasida ham `ValueError: Unable to read
workbook: ...` beradi, ya'ni aniq tipni ushlash yo'li YO'Q.

Keng ushlash O'Z kodimizdagi xatoni yashirmasligi uchun `try` bloki
ATAYIN TOR — unda faqat kutubxona chaqiruvlari turadi, bizning
mantiqimiz esa tashqarida. Xom istisno to'liq log'ga yoziladi.
"""


def read_rows(
    raw: bytes,
    *,
    expected_columns: int,
    limits: ReadLimits | None = None,
) -> list[SheetRow]:
    """Yuklangan `.xlsx` dan ma'lumot qatorlarini o'qiydi.

    Ketma-ketlik MAJBURIY (modul docstringi): hajm -> ZIP -> parse ->
    varaq/qator/ustun. Har bosqich `ImportRejected` ko'taradi.

    Args:
        raw: yuklangan faylning XOM baytlari (ishonchsiz).
        expected_columns: shablon ustunlari soni; qisqa qator shu songacha
            `None` bilan to'ldiriladi (qisqa qator XATO EMAS — validator
            unga `row_too_short` beradi va foydalanuvchi qator raqamini
            ko'radi).
        limits: chegaralar; `None` bo'lsa modul standartlari.

    Raises:
        ImportRejected: `file_too_large` / `file_too_complex` /
            `unsupported_file_type`.
    """
    gates = ReadLimits() if limits is None else limits

    # --- 1. Xom hajm. Bu darvoza faylni XOTIRAGA olishdan oldin emas,
    # olgandan KEYIN ishlaydi (router `await file.read()` qiladi), lekin
    # baribir birinchi: usiz 4 GB fayl ZIP darvozasigacha yetib borardi.
    if len(raw) > gates.upload_bytes:
        raise ImportRejected(_FILE_TOO_LARGE)

    _check_zip(raw, gates)

    workbook = _load(raw)
    try:
        if len(workbook.sheetnames) > gates.sheets:
            raise ImportRejected(_FILE_TOO_COMPLEX)
        return list(_iter_rows(workbook, gates, expected_columns))
    finally:
        # `read_only=True` fayl deskriptorlarini OCHIQ qoldiradi —
        # `close()` siz ular so'rovlar bo'ylab to'planardi.
        workbook.close()


def _check_zip(raw: bytes, gates: ReadLimits) -> None:
    """ZIP darvozasi — `load_workbook()` DAN OLDIN (T-02-88).

    OCHILGAN hajm `zf.infolist()` dagi e'lon qilingan `file_size` lar
    yig'indisi bo'yicha o'lchanadi. Bu qiymatni HUJUMCHI yozadi, lekin
    CPython `zipfile` uni O'ZI ham chegara sifatida majburlaydi:
    `ZipExtFile` `file_size` dan ortiq bayt QAYTARMAYDI. Ya'ni yolg'on
    kichik `file_size` yozgan arxiv bizni aldab o'tolmaydi — u shunchaki
    kesilgan ma'lumot beradi va parse xatosiga tushadi.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        # `.xlsx` bo'lmagan har qanday fayl (PDF, .xls, .csv, rasm) shu
        # yerda to'xtaydi — kengaytma tekshiruvidan MUSTAQIL ravishda.
        log.info("import_not_a_zip", error=str(exc))
        raise ImportRejected(_UNSUPPORTED) from exc

    if sum(entry.file_size for entry in entries) > gates.uncompressed_bytes:
        log.warning("import_zip_bomb_rejected", entries=len(entries))
        raise ImportRejected(_FILE_TOO_LARGE)

    if len(entries) > gates.zip_entries:
        raise ImportRejected(_FILE_TOO_COMPLEX)


def _load(raw: bytes) -> Any:
    """`openpyxl.load_workbook` — `read_only` va `data_only` IKKALASI ham.

    `read_only=True` — lazy yuklash, doimiy xotira (T-02-89).
    `data_only=True` — formula MATNI emas, hisoblangan QIYMAT. Usiz
    `=SUM(A1:A9)` yozilgan katak validatorga formula satri bo'lib
    kelardi va "rasta kodi noto'g'ri" degan chalg'ituvchi xato berardi.

    Chaqiruv `openpyxl.load_workbook` ORQALI (modul atributi sifatida),
    to'g'ridan-to'g'ri import qilingan nom bilan EMAS: `test_zip_bomb_is_
    rejected_before_parsing` uni monkeypatch bilan almashtirib, ZIP
    darvozasi ishlaganda bu funksiya CHAQIRILMAGANINI tekshiradi.
    """
    try:
        return openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except _PARSE_FAILURES as exc:
        log.warning("import_parse_failed", error=str(exc), kind=type(exc).__name__)
        raise ImportRejected(_UNSUPPORTED) from exc


def _iter_rows(workbook: Any, gates: ReadLimits, expected_columns: int) -> Iterator[SheetRow]:
    """Birinchi varaqni o'qiydi va `SheetRow` oqimini beradi.

    ⚠ CHEGARA IKKI MARTA TEKSHIRILADI VA BU ATAYIN.

    `ws.max_row` / `ws.max_column` faylning O'ZIDAGI `<dimension>`
    e'lonidan keladi, ya'ni ularni HUJUMCHI yozadi. Ikkita mustaqil
    yolg'on shakli bor va ikkalasi ham o'lchandi:

    * e'lon KATTA — arzon rad etish (birinchi tekshiruv);
    * e'lon YO'Q — o'shanda `max_row` `None` bo'ladi va birinchi
      tekshiruv JIMGINA o'tib ketadi. Yagona qoladigan darvoza —
      iteratsiya paytidagi sanoq.

    (E'lon KICHIK bo'lgan uchinchi holat DoS emas: openpyxl o'sha
    e'londan ortiq qator qaytarmaydi, ya'ni hujumchi faqat O'Z
    ma'lumotini yo'qotadi — o'lchandi.)
    """
    worksheet = workbook[workbook.sheetnames[0]]

    declared_rows = worksheet.max_row or 0
    declared_cols = worksheet.max_column or 0
    if declared_rows > gates.rows or declared_cols > gates.cols:
        raise ImportRejected(_FILE_TOO_COMPLEX)

    seen = 0
    for number, cells in enumerate(_safe_rows(worksheet), start=HEADER_ROW):
        seen += 1
        if seen > gates.rows or len(cells) > gates.cols:
            raise ImportRejected(_FILE_TOO_COMPLEX)
        if number == HEADER_ROW:
            # Sarlavha MATNI o'qilmaydi (O-05) — faqat tashlanadi.
            continue

        values = tuple(_cell_text(cell) for cell in cells[:expected_columns])
        if not any(values):
            # BUTUNLAY BO'SH QATOR TASHLANADI. Excel'dan eksport qilingan
            # faylda formatlangan, lekin bo'sh qatorlar odatiy hol; ularni
            # saqlash har biriga `empty_code` xatosi berardi va D-14
            # (all-or-nothing) tufayli MUTLAQO TO'G'RI fayl ham butunlay
            # rad etilardi.
            continue

        yield SheetRow(row=number, values=_pad(values, expected_columns))


def _safe_rows(worksheet: Any) -> Iterator[tuple[object, ...]]:
    """`iter_rows()` ni LAZY o'raydi va parse xatosini `ImportRejected` ga o'giradi.

    ⚠ NATIJA RO'YXATGA YIG'ILMAYDI (`list(...)` YO'Q) va bu darvoza
    uchun HAL QILUVCHI. `<dimension>` e'loni bo'lmagan faylda openpyxl
    varaqni oxirigacha o'qiydi; qatorlarni avval ro'yxatga yig'ib,
    keyin sanash "chegara bor" degan yolg'on tuyg'u berardi — xotira
    o'sha paytda ALLAQACHON yeyilgan bo'lardi (T-02-89).

    `StopIteration` QO'LDA ushlanadi: generator ichida u sizib chiqsa
    Python uni `RuntimeError` ga aylantiradi (PEP 479) va butun so'rov
    500 bilan tugardi.
    """
    rows = worksheet.iter_rows(values_only=True)
    while True:
        try:
            row = next(rows)
        except StopIteration:
            return
        except _PARSE_FAILURES as exc:
            log.warning("import_rows_parse_failed", error=str(exc), kind=type(exc).__name__)
            raise ImportRejected(_UNSUPPORTED) from exc
        yield row


def _pad(values: tuple[str | None, ...], width: int) -> tuple[str | None, ...]:
    """Qatorni `width` gacha `None` bilan to'ldiradi (qisqa qator xato EMAS)."""
    if len(values) >= width:
        return values
    return values + (None,) * (width - len(values))


def _cell_text(cell: object) -> str | None:
    """Katak qiymatini matnga keltiradi; bo'sh katak -> `None`.

    Uchta holat ATAYIN alohida ishlanadi va uchalasi ham real fayllardan
    keladi:

    * `datetime`/`date` — Excel sanani sana TIPIDA saqlaydi. Oddiy
      `str()` `"2026-08-01 00:00:00"` berardi va `date.fromisoformat()`
      uni RAD ETARDI, ya'ni to'g'ri to'ldirilgan fayl `invalid_date`
      olardi;
    * butun qiymatli `float` — Excel `12` ni ba'zan `12.0` qilib
      qaytaradi. `str()` `"12.0"` berardi va rasta kodi `"12"` bilan
      MOS TUSHMASDI (D-15 idempotentligi buzilardi: har import yangi
      kod yaratardi);
    * `bool` — `str()` inglizcha `"True"` beradi. Bu shablonda
      ishlatilmaydi, lekin kimdir ustunni belgilab qo'yishi mumkin,
      shuning uchun aniq matnga keltiriladi.
    """
    if cell is None:
        return None
    if isinstance(cell, bool):
        return "true" if cell else "false"
    if isinstance(cell, datetime):
        text = cell.date().isoformat()
    elif isinstance(cell, date):
        text = cell.isoformat()
    elif isinstance(cell, float) and cell.is_integer():
        text = str(int(cell))
    else:
        text = str(cell)

    stripped = text.strip()
    return stripped or None
